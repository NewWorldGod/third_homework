# -*- coding: utf-8 -*-
"""向报告追加第8章：进阶任务（选做）——重复惩罚生成对比"""
import docx
from docx.shared import Pt

OUT = r"d:\作业\大模型作业三\2025310291_杨锐_实验作业三.docx"
doc = docx.Document(OUT)

style = doc.tables[13].style  # 沿用模板表格样式


def add_para(text, bold=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    return p


# 章节标题与说明
add_para("8  进阶任务（选做）：重复惩罚生成对比", bold=True)
add_para(
    "任务四(b)：给 generate() 增加重复惩罚并对比前后生成效果。实现见 "
    "hw3_min_llm/repetition_penalty.py（复用基线 checkpoint），共两种惩罚机制：")
add_para(
    "① 频率惩罚（软惩罚）：对已生成过的字符，按其出现次数 n 对 logits 减去 alpha×n，"
    "使高频字符越来越难被再次采样（本实验取 alpha=1.0 与 1.5）；")
add_para(
    "② 3-gram 禁止（硬惩罚）：维护“已出现 3 元组”集合，若候选字符会与前面两字构成"
    "已出现过的 3 元组，则将其 logits 置 -inf 直接禁止；若某上下文下候选全被禁止则放弃"
    "惩罚以保证不卡死。对比实验固定基线模型、top_k=20，在 temperature=0.5 与 1.0 下"
    "用提示词「春」「月」各生成，统计“重复 3-gram 比例”与“字符多样性（去重字符数/总字符数）”"
    "两个指标，完整原文见 repetition_penalty_results.txt。")

# 结果表
add_para("表 7  重复惩罚前后生成对比（提示词「月」，120 字）", bold=True)
t = doc.add_table(rows=6, cols=4)
t.style = style
hdr = ["设置", "温度", "重复 3-gram 比例", "字符多样性"]
rows = [
    ("原始 generate()", "0.5", "0.8%", "67.8%"),
    ("原始 generate()", "1.0", "0.8%", "74.4%"),
    ("频率惩罚 alpha=1.5", "1.0", "0.0%", "79.3%"),
    ("3-gram 禁止", "0.5", "0.0%", "71.9%"),
    ("3-gram 禁止 + 频率惩罚 1.0", "1.0", "0.0%", "76.9%"),
]
for j, h in enumerate(hdr):
    c = t.rows[0].cells[j]
    c.paragraphs[0].add_run(h).bold = True
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        t.rows[i + 1].cells[j].paragraphs[0].add_run(v).font.size = Pt(9)

add_para("")

# 前后生成原文对比
add_para("前后生成原文对比（temperature=1.0，提示词「月」开头）：", bold=True)
add_para(
    "惩罚前（原始 generate()）：\n"
    "月出惊山鸟，时花落木萧萧下，不尽，不尽长江滚滚来。\n"
    "花近高楼伤客心，万方多难此登临。锦江春色来天地，玉垒浮云变古今。……", size=9)
add_para(
    "惩罚后（3-gram 禁止，同一随机种子）：\n"
    "月出惊山鸟，时花落木萧萧下，不尽，不复长江滚滚来。\n"
    "花近高楼伤客心，万方多难此登临。锦江春色来天地，玉垒浮云变古今。……", size=9)
add_para(
    "惩罚后（频率惩罚 alpha=1.5，提示词「春」节选，展示副作用）：\n"
    "辛苦遭逢起一经知干戈寥落四周星行成三月之悠纷纷未啼不与君莫笑农家腊酒浑，"
    "丰年留客足鸡豚。……", size=9)

add_para("结果分析：", bold=True)
add_para(
    "① 重复来源：本模型的“重复”主要是整句背诵式复读（如“不尽，不尽”相邻复现），"
    "重复 3-gram 比例本身不高（0.8%），因为语料被背下后多数 3 元组本就只按原句顺序出现。"
    "② 3-gram 禁止效果最稳：把重复率降为 0，字符多样性提升（67.8%→71.9%），且被禁后"
    "模型会改采次优字符（“不尽”变“不复”），只在被禁止处产生局部扰动，整句背诵格式"
    "基本保留。③ 频率惩罚是一把双刃剑：alpha=1.5 时确实提高了多样性（74.4%→79.3%），"
    "但惩罚随出现次数累积，生成后半段高频常用字（的、不、来等）被过度压制，导致"
    "“起一经知干戈寥落四周星行成三月之悠纷纷”这类跨句乱序拼接——与采样分析中"
    "“重复↔胡言”的权衡一致：惩罚过强等于把分布推向“胡言”一端。④ 组合使用"
    "（3-gram 禁止 + 轻频率惩罚 1.0）兼顾两者：重复率 0、多样性 76.9%，流畅性损失"
    "可控。结论：对以背诵为主的小模型，硬性 3-gram 禁止是性价比最高的重复惩罚；"
    "频率惩罚的 alpha 必须保守设置。")

doc.save(OUT)
print("appended section 8, saved:", OUT)
